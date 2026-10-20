"""
RAG Engine สำหรับแชทบอทการจัดการขยะและการรีไซเคิล
Flow:  (ภาพ -> ระบุประเภทเบื้องต้น) -> ค้นคืนข้อมูล (embedding + cosine) -> LLM ตอบพร้อมอ้างอิง [n]
"""
import base64
import hashlib
import importlib.util
import json
import os
import math
import re
from collections import Counter
from pathlib import Path

import numpy as np
import storage
from dotenv import load_dotenv

load_dotenv()

BASE_DIR = Path(__file__).parent
DATA_DIR = storage.DATA_DIR
CACHE_DIR = storage.CACHE_DIR
AREA_NAME = os.getenv("AREA_NAME", "").strip()  # เช่น "เทศบาลเมือง..." ใช้แนะนำหน่วยงานที่ติดต่อ
EMBED_MODEL = "paraphrase-multilingual-MiniLM-L12-v2"  # รองรับภาษาไทย

SYSTEM_PROMPT = """คุณคือ "น้องรีไซเคิล" ผู้ช่วย AI ด้านการคัดแยกขยะและการรีไซเคิล

กติกา:
1. ตอบจากข้อมูลใน <context> เท่านั้น ห้ามแต่งข้อมูลที่ไม่มีในบริบท
2. ทุกประโยคที่ใช้ข้อมูลจากบริบท ให้ใส่หมายเลขอ้างอิงท้ายประโยค เช่น [1] หรือ [2][3] ตามหมายเลขของบริบท
3. ถ้าบริบทไม่พอ ให้บอกตรง ๆ ว่าไม่มีข้อมูลเพียงพอ และแนะนำให้สอบถามเทศบาล/สำนักงานเขตในพื้นที่
4. ถ้ามี <image_analysis> แสดงว่าเป็นผลระบุประเภทจากภาพถ่ายซึ่งเป็นการประมาณเบื้องต้น ให้บอกผู้ใช้ว่าอาจคลาดเคลื่อน
   และถ้าความมั่นใจต่ำหรือไม่ชัดเจน ให้ขอภาพ/รายละเอียดเพิ่ม
5. เมื่อผู้ใช้ถามเรื่องการทิ้ง จัดเก็บ หรือเตรียมขยะ (หรือส่งภาพขยะมา) ให้ตอบตามโครงสร้างนี้ (ข้ามหัวข้อที่ไม่มีข้อมูล):
   **ประเภทขยะ / ถังที่ควรทิ้ง**
   **วิธีเตรียมก่อนทิ้งหรือรีไซเคิล** (เป็นข้อ ๆ)
   **วิธีจัดเก็บ**
   **ข้อควรระวัง**
6. ตอบเป็นภาษาไทย กระชับ เป็นมิตร และเน้นความปลอดภัยเสมอสำหรับขยะอันตราย
7. ถ้าคำตอบระบุว่าขยะควรทิ้งถังใด ให้ขึ้นต้นคำตอบด้วยบรรทัดเดียวในรูปแบบ  BIN: สี  โดยสีต้องเป็น เขียว, เหลือง, น้ำเงิน หรือ แดง เท่านั้น
   (หลายถังคั่นด้วยจุลภาค เช่น BIN: เหลือง,น้ำเงิน เมื่อต้องแยกส่วนทิ้งคนละถัง) แล้วจึงขึ้นบรรทัดใหม่เขียนคำตอบตามปกติ
   สีต้องอิงจากบริบทเท่านั้น ถ้าไม่เกี่ยวกับถังขยะหรือบริบทไม่ระบุ ห้ามใส่บรรทัดนี้
   บรรทัด BIN ต้องเป็นข้อความธรรมดา ไม่ใส่ตัวหนา เครื่องหมาย หรือคำว่า "สี" นำหน้าชื่อสี
8. ถ้าบริบทมีจุดรับพร้อมลิงก์แผนที่ ให้แนะนำชื่อ เวลาทำการ และลิงก์แผนที่ตามที่ปรากฏในบริบทเท่านั้น ห้ามแต่งสถานที่หรือเวลาที่ไม่มีในบริบท"""

VISION_PROMPT = """คุณช่วยระบุประเภทขยะจากภาพถ่ายเพื่อใช้ค้นข้อมูลการคัดแยก
ตอบเป็น JSON เท่านั้น ห้ามมีข้อความอื่น รูปแบบ:
{"items":[{"name":"ชื่อสิ่งของภาษาไทย","material":"วัสดุ เช่น พลาสติก PET/กระดาษ/แก้ว/อะลูมิเนียม/แบตเตอรี่","condition":"สภาพ เช่น สะอาด/เปื้อนอาหาร/มีของเหลือ"}],
"confidence":"high|medium|low","note":"ข้อสังเกตสั้น ๆ"}
ถ้าภาพไม่ใช่ขยะ หรือมองไม่ชัด ให้ items เป็น [] และอธิบายใน note"""


# ---------- การเตรียมเอกสาร ----------
SUPPORTED_EXT = {".md", ".txt", ".pdf", ".docx"}

_MARKS = "\u0e31\u0e34-\u0e3a\u0e47-\u0e4e"  # สระบน/ล่าง วรรณยุกต์ ฯลฯ


def clean_thai(text: str) -> str:
    """แก้ปัญหาภาษาไทยที่พบบ่อยจากการดึงข้อความ PDF (สระ/วรรณยุกต์ลอย, เว้นวรรคแทรก, ำ แตกเป็น 2 ตัว)"""
    # "ํ" + "า" -> "ำ" (รวมกรณีมีวรรณยุกต์คั่น)
    text = re.sub("\u0e4d([\u0e48-\u0e4b])\u0e32", "\\1\u0e33", text)
    text = text.replace("\u0e4d\u0e32", "\u0e33")
    # วรรณยุกต์มาก่อนสระบน -> สลับให้ถูกลำดับ (สระ แล้วค่อยวรรณยุกต์)
    text = re.sub("([\u0e48-\u0e4b])([\u0e31\u0e34-\u0e37\u0e47])", "\\2\\1", text)
    # เว้นวรรคนำหน้าสระ/วรรณยุกต์ ไม่มีในภาษาไทย -> ลบ
    text = re.sub(f"[ \t]+([{_MARKS}])", "\\1", text)
    # สระหน้า (เ แ โ ใ ไ) ตามด้วยเว้นวรรคแล้วพยัญชนะ -> ลบเว้นวรรค
    text = re.sub("([เแโใไ])[ \t]+(?=[\u0e01-\u0e2e])", "\\1", text)
    # บรรทัดที่ถูกเว้นวรรคระหว่างตัวอักษรทุกตัว (เช่น "ก า ร จ ั ด") -> รวมกลับ
    fixed = []
    for line in text.split("\n"):
        tokens = line.split(" ")
        if (len(tokens) >= 6 and re.search("[\u0e01-\u0e2e]", line)
                and sum(len(t) == 1 for t in tokens) / len(tokens) > 0.6):
            line = "".join(tokens)
        fixed.append(line)
    return "\n".join(fixed)


def extract_text(path: Path) -> str:
    """อ่านข้อความจาก .md/.txt/.pdf/.docx (PDF แยกตามหน้า, DOCX แปลง Heading เป็น ##)"""
    ext = path.suffix.lower()
    if ext in {".md", ".txt"}:
        return path.read_text(encoding="utf-8-sig", errors="ignore")
    if ext == ".pdf":
        try:  # PyMuPDF อ่านภาษาไทยได้ดีกว่า pypdf
            import pymupdf
            with pymupdf.open(str(path)) as pdf:
                raw_pages = [pg.get_text() for pg in pdf]
        except ImportError:
            from pypdf import PdfReader
            raw_pages = [pg.extract_text() or "" for pg in PdfReader(str(path)).pages]
        pages = []
        for i, raw in enumerate(raw_pages, 1):
            t = clean_thai(raw).strip()
            if t:
                pages.append(f"# หน้า {i}\n{t}")
        return "\n\n".join(pages)
    if ext == ".docx":
        from docx import Document
        doc = Document(str(path))
        lines = []
        for para in doc.paragraphs:
            t = para.text.strip()
            if not t:
                continue
            style = (para.style.name if para.style is not None else "").lower()
            lines.append(f"## {t}" if style.startswith(("heading", "title")) else t)
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines)
    return ""


def _paragraphs(body: str, max_chars: int):
    """แบ่งย่อหน้า ถ้ายาวเกิน (เช่น PDF ที่ขึ้นบรรทัดใหม่ทุกบรรทัด/ภาษาไทยไม่มีเว้นวรรค) ให้รวม/ตัดตามความยาว"""
    for para in re.split(r"\n\s*\n|\n(?=[-*•] )", body):
        para = para.strip()
        if len(para) <= max_chars:
            yield para
            continue
        buf = ""
        for line in para.split("\n"):
            line = line.strip()
            while len(line) > max_chars:
                if buf:
                    yield buf
                    buf = ""
                yield line[:max_chars]
                line = line[max_chars:]
            if buf and len(buf) + len(line) + 1 > max_chars:
                yield buf
                buf = ""
            buf = f"{buf}\n{line}" if buf else line
        if buf:
            yield buf

def chunk_text(text: str, source: str, max_chars: int = 700):
    """แบ่งตามหัวข้อ (#) ก่อน แล้วค่อยแบ่งตามย่อหน้าถ้ายาวเกิน"""
    chunks = []
    for sec in re.split(r"\n(?=#{1,3} )", text):
        sec = sec.strip()
        if not sec:
            continue
        heading = sec.split("\n")[0] if sec.startswith("#") else ""
        title = heading.lstrip("# ").strip()
        if len(sec) <= max_chars:
            chunks.append({"text": sec, "source": source, "title": title})
            continue
        buf = heading
        for para in _paragraphs(sec[len(heading):] if heading else sec, max_chars):
            para = para.strip()
            if not para or para == heading:
                continue
            if len(buf) + len(para) > max_chars and buf.strip() != heading:
                chunks.append({"text": buf.strip(), "source": source, "title": title})
                buf = heading
            buf += "\n" + para
        if buf.strip() and buf.strip() != heading:
            chunks.append({"text": buf.strip(), "source": source, "title": title})
    return chunks


def load_documents():
    """คืน (chunks, skipped) — skipped คือไฟล์ที่อ่านไม่ได้/ไม่มีข้อความ"""
    chunks, skipped = [], []
    for path in sorted(DATA_DIR.glob("*")):
        if path.suffix.lower() not in SUPPORTED_EXT:
            continue
        try:
            text = extract_text(path)
        except Exception as e:  # ไฟล์เสีย/มีรหัสผ่าน ฯลฯ
            skipped.append((path.name, f"อ่านไฟล์ไม่ได้: {e}"))
            continue
        if not text.strip():
            skipped.append((path.name, "ไม่พบข้อความ (อาจเป็น PDF สแกน/รูปภาพ)"))
            continue
        chunks += chunk_text(text, path.name)
    try:  # จุดรับขยะจาก locations/points.csv (จัดการที่หน้า Admin)
        from places import load_places, place_chunk
        for p in load_places():
            chunks.append({"text": place_chunk(p), "source": "จุดรับขยะ", "title": f"{p['type']}: {p['name']}"})
    except Exception as e:
        skipped.append(("locations/points.csv", f"อ่านไม่ได้: {e}"))
    return chunks, skipped


def format_vision(vision: dict) -> str:
    if not vision or not vision.get("items"):
        return ""
    lines = []
    for it in vision["items"]:
        parts = [it.get("name", "")]
        if it.get("material"):
            parts.append(f"วัสดุ: {it['material']}")
        if it.get("condition"):
            parts.append(f"สภาพ: {it['condition']}")
        lines.append("- " + " | ".join(p for p in parts if p))
    return "\n".join(lines)


# แก้คำสะกดผิด/คำสะกดอื่นที่พบบ่อยในคำถาม ก่อนค้นหา (เพิ่มได้ตามต้องการ)
QUERY_FIXES = [
    (r"กะป๋อง", "กระป๋อง"),
    (r"สเปรย(?!์)", "สเปรย์"),
    (r"สเปร์", "สเปรย์"),
    (r"รีไซเคิ[ล้่ลน]+|รีไซเคิ้ล", "รีไซเคิล"),
    (r"พาสติก|พลาสติค|พลาสติ๊ก", "พลาสติก"),
    (r"อลูมิเนียม|อะลูมิเนี่ยม", "อะลูมิเนียม"),
    (r"อิเล็กทรอนิค(?!ส์)", "อิเล็กทรอนิกส์"),
]
_PUNCT = re.compile(r"[\s.,;:!?()\[\]{}\"'“”‘’\-–—/\\|*#<>_]+")


def fix_query(q: str) -> str:
    for pat, repl in QUERY_FIXES:
        q = re.sub(pat, repl, q)
    return q


def _ngrams(text: str, n: int = 3) -> set:
    """ชุดตัวอักษรต่อเนื่อง n ตัว — ใช้จับคำตรงกัน โดยไม่ต้องตัดคำไทย และทนต่อคำสะกดผิดเล็กน้อย"""
    t = _PUNCT.sub("", text.lower())
    return {t[i:i + n] for i in range(len(t) - n + 1)} if len(t) >= n else ({t} if t else set())


class OpenAIEmbedder:
    """embedding ผ่าน API: ไม่ต้องโหลด torch/โมเดลในเครื่อง เริ่มระบบได้ทันที"""
    DIMS = {"text-embedding-3-small": 1536, "text-embedding-3-large": 3072, "text-embedding-ada-002": 1536}

    def __init__(self):
        from openai import OpenAI
        self.client = OpenAI()
        self.model = os.getenv("OPENAI_EMBED_MODEL", "text-embedding-3-small")

    def get_sentence_embedding_dimension(self):
        return self.DIMS.get(self.model, 1536)

    def encode(self, texts, normalize_embeddings=True, show_progress_bar=False):
        texts = list(texts)
        if not texts:
            return np.zeros((0, self.get_sentence_embedding_dimension()), dtype=np.float32)
        vecs = []
        for i in range(0, len(texts), 96):
            resp = self.client.embeddings.create(model=self.model, input=texts[i:i + 96])
            vecs += [d.embedding for d in resp.data]
        arr = np.array(vecs, dtype=np.float32)
        if normalize_embeddings:
            arr /= np.linalg.norm(arr, axis=1, keepdims=True)
        return arr


def make_embedder():
    """EMBEDDINGS=openai: ใช้ API (เริ่มเร็ว ไม่ใช้ torch) | local: โมเดลในเครื่อง (ฟรี ใช้ออฟไลน์ได้)
    auto (ค่าเริ่มต้น): ใช้โมเดลในเครื่องถ้าติดตั้งไว้ ไม่เช่นนั้นใช้ openai"""
    mode = os.getenv("EMBEDDINGS", "auto").lower()
    have_key = bool(os.getenv("OPENAI_API_KEY"))
    have_local = importlib.util.find_spec("sentence_transformers") is not None
    if mode == "openai" and not have_key:
        print("EMBEDDINGS=openai แต่ไม่มี OPENAI_API_KEY — ลองใช้โมเดลในเครื่องแทน")
    if have_key and (mode == "openai" or (mode == "auto" and not have_local)):
        e = OpenAIEmbedder()
        return e, f"openai:{e.model}"
    if not have_local:
        raise RuntimeError("ไม่พบ sentence-transformers สำหรับค้นหาในเครื่อง: ติดตั้งด้วย "
                           "pip install -r requirements-local.txt หรือตั้ง EMBEDDINGS=openai พร้อม OPENAI_API_KEY")
    from sentence_transformers import SentenceTransformer  # import ช้า (torch) จึงโหลดตอนนี้
    try:  # ใช้ไฟล์โมเดลในเครื่องก่อน (เร็ว ไม่ต้องเช็คอินเทอร์เน็ต) ถ้ายังไม่มีค่อยดาวน์โหลด
        model = SentenceTransformer(EMBED_MODEL, local_files_only=True)
    except Exception:
        model = SentenceTransformer(EMBED_MODEL)
    return model, EMBED_MODEL


class WasteRAG:
    def __init__(self):
        self.embedder, self.embed_id = make_embedder()
        DATA_DIR.mkdir(exist_ok=True)
        self.chunks, self.skipped = load_documents()
        self.embeddings = self._load_or_build_embeddings()
        if self.embed_id == EMBED_MODEL and self.chunks:  # อุ่นเครื่อง ให้คำถามแรกไม่ช้า
            self.embedder.encode(["อุ่นเครื่อง"], normalize_embeddings=True, show_progress_bar=False)

        # เลือกผู้ให้บริการ LLM: LLM_PROVIDER=openai | anthropic (ถ้าไม่ระบุ จะดูจาก API key ที่มี)
        self.client, self.provider, self.model = None, None, None
        provider = os.getenv("LLM_PROVIDER", "").lower()
        if not provider:
            provider = "openai" if (os.getenv("OPENAI_API_KEY") and not os.getenv("ANTHROPIC_API_KEY")) else "anthropic"
        if provider == "openai" and os.getenv("OPENAI_API_KEY"):
            from openai import OpenAI
            self.client, self.provider = OpenAI(), "openai"
            self.model = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
        elif provider == "anthropic" and os.getenv("ANTHROPIC_API_KEY"):
            import anthropic
            self.client, self.provider = anthropic.Anthropic(), "anthropic"
            self.model = os.getenv("ANTHROPIC_MODEL", "claude-sonnet-5-5")

    # ---------- LLM wrappers (รองรับ OpenAI / Anthropic) ----------
    def _complete(self, system: str, messages: list, max_tokens: int) -> str:
        if self.provider == "openai":
            resp = self.client.chat.completions.create(
                model=self.model,
                max_completion_tokens=max_tokens,
                messages=[{"role": "system", "content": system}] + messages,
            )
            return resp.choices[0].message.content or ""
        resp = self.client.messages.create(model=self.model, max_tokens=max_tokens,
                                           system=system, messages=messages)
        return resp.content[0].text

    def _vision(self, image_bytes: bytes, media_type: str) -> str:
        b64 = base64.standard_b64encode(image_bytes).decode()
        if self.provider == "openai":
            resp = self.client.chat.completions.create(
                model=self.model,
                max_completion_tokens=500,
                messages=[
                    {"role": "system", "content": VISION_PROMPT},
                    {"role": "user", "content": [
                        {"type": "image_url", "image_url": {"url": f"data:{media_type};base64,{b64}"}},
                        {"type": "text", "text": "ระบุขยะในภาพนี้"},
                    ]},
                ],
            )
            return resp.choices[0].message.content or ""
        resp = self.client.messages.create(
            model=self.model, max_tokens=500, system=VISION_PROMPT,
            messages=[{"role": "user", "content": [
                {"type": "image", "source": {"type": "base64", "media_type": media_type, "data": b64}},
                {"type": "text", "text": "ระบุขยะในภาพนี้"},
            ]}],
        )
        return resp.content[0].text

    def _load_or_build_embeddings(self):
        if not self.chunks:
            return np.zeros((0, self.embedder.get_sentence_embedding_dimension()))
        CACHE_DIR.mkdir(exist_ok=True)
        digest = hashlib.md5((self.embed_id + "".join(c["text"] for c in self.chunks)).encode()).hexdigest()
        cache_file = CACHE_DIR / f"emb_{digest}.npy"
        if cache_file.exists():
            return np.load(cache_file)
        emb = self.embedder.encode([c["text"] for c in self.chunks],
                                   normalize_embeddings=True, show_progress_bar=False)
        for old in CACHE_DIR.glob("emb_*.npy"):  # ล้างแคชเก่า
            old.unlink(missing_ok=True)
        np.save(cache_file, emb)
        return emb

    def reload(self):
        """อ่านเอกสารใน data/ ใหม่และอัปเดตดัชนี (ไม่ต้องรีสตาร์ทแอป)"""
        self.chunks, self.skipped = load_documents()
        self.embeddings = self._load_or_build_embeddings()

    # ---------- Retrieval ----------
    def _ensure_keyword_index(self):
        if getattr(self, "_kw_for", None) is self.chunks:
            return
        self.chunk_grams = [_ngrams(c["text"]) for c in self.chunks]
        df = Counter(g for grams in self.chunk_grams for g in grams)
        n = len(self.chunks)
        self.idf = {g: math.log((n + 1) / (c + 0.5)) + 0.1 for g, c in df.items()}  # คำที่พบน้อย = สำคัญมาก
        self.idf_unknown = math.log((n + 1) / 0.5) + 0.1
        self._kw_for = self.chunks

    def retrieve(self, query: str, k: int = 4, min_score: float = 0.25):
        """ค้นแบบผสม: ความหมาย (embedding) + คำตรงกัน (n-gram ถ่วงน้ำหนักคำหายาก)
        ใช้ข้อมูลชิ้นนั้นเมื่อ ความหมาย >= min_score หรือ คำตรง >= 0.5"""
        if not self.chunks:
            return []
        self._ensure_keyword_index()
        query = fix_query(query)
        q = self.embedder.encode([query], normalize_embeddings=True)[0]
        sem = self.embeddings @ q  # cosine similarity (normalize แล้ว)
        qg = _ngrams(query)
        denom = sum(self.idf.get(g, self.idf_unknown) for g in qg) or 1.0
        kw = np.array([sum(self.idf.get(g, 0.0) for g in qg & cg) / denom for cg in self.chunk_grams])
        score = sem + 0.5 * kw
        hits = []
        for i in np.argsort(-score)[:k]:
            if sem[i] >= min_score or kw[i] >= 0.5:
                hits.append({**self.chunks[i], "score": float(min(score[i], 1.0)),
                             "semantic": float(sem[i]), "keyword": float(kw[i])})
        for n, h in enumerate(hits, 1):
            h["ref"] = n
        return hits

    # ---------- Image-assisted query ----------
    def identify_image(self, image_bytes: bytes, media_type: str = "image/jpeg") -> dict:
        if self.client is None:
            raise RuntimeError("ฟีเจอร์วิเคราะห์ภาพต้องตั้งค่า API key (OPENAI_API_KEY หรือ ANTHROPIC_API_KEY) ในไฟล์ .env")
        raw = self._vision(image_bytes, media_type)
        m = re.search(r"\{.*\}", raw, re.S)
        try:
            data = json.loads(m.group(0)) if m else {}
        except json.JSONDecodeError:
            data = {}
        data.setdefault("items", [])
        data.setdefault("confidence", "low")
        data.setdefault("note", "" if data["items"] else raw.strip()[:200])
        data["query"] = " ".join(
            f"{i.get('name', '')} {i.get('material', '')} {i.get('condition', '')}" for i in data["items"]
        ).strip()
        return data

    # ---------- Answer ----------
    def _prepare(self, question, history, image_bytes, k):
        """ค้นหา (+วิเคราะห์ภาพ) แล้วเตรียมข้อความส่ง LLM
        คืน (hits, vision, system, msgs, immediate) — immediate ไม่ว่าง = ตอบได้เลยโดยไม่ต้องเรียก LLM"""
        history = history or []
        vision, query = None, question

        if image_bytes:
            vision = self.identify_image(image_bytes)
            query = f"{question} {vision['query']}".strip()
        elif len(question) < 15:  # คำถามสั้นมาก (เช่น "แล้วฝาล่ะ") ใช้คำถามก่อนหน้าช่วยค้น
            prev = [m["content"] for m in history if m["role"] == "user"]
            if prev:
                query = f"{prev[-1]} {question}"

        hits = self.retrieve(query, k=k)

        if not hits:
            msg = "ขออภัย ไม่พบข้อมูลที่เกี่ยวข้องในฐานความรู้"
            if vision is not None and not vision["items"]:
                msg += f" และระบุขยะในภาพไม่ได้ ({vision.get('note', '')}) ลองถ่ายภาพให้ชัดขึ้นหรือพิมพ์อธิบายเพิ่มนะ"
            else:
                msg += " ลองถามให้เฉพาะเจาะจงขึ้น หรือสอบถามหน่วยงานท้องถิ่นโดยตรง"
            return [], vision, None, None, msg

        if self.client is None:  # โหมดไม่มี API key
            body = "\n\n---\n\n".join(f"[{h['ref']}] {h['text']}" for h in hits)
            return hits, vision, None, None, "(โหมดไม่มี LLM — แสดงข้อมูลที่ค้นพบ)\n\n" + body

        context = "\n\n".join(f"[{h['ref']}] (แหล่งที่มา: {h['source']})\n{h['text']}" for h in hits)
        image_block = ""
        if vision:
            image_block = (f"<image_analysis>\nความมั่นใจ: {vision['confidence']}\n"
                           f"{format_vision(vision) or 'ระบุไม่ได้'}\nหมายเหตุ: {vision.get('note', '')}\n"
                           f"</image_analysis>\n\n")

        msgs = [{"role": m["role"], "content": m["content"]} for m in history[-6:]]
        while msgs and msgs[0]["role"] != "user":
            msgs.pop(0)
        msgs.append({"role": "user",
                     "content": f"<context>\n{context}\n</context>\n\n{image_block}คำถาม: {question}"})

        system = SYSTEM_PROMPT
        if AREA_NAME:
            system += (f"\nพื้นที่ให้บริการ: {AREA_NAME} "
                       "ถ้าบริบทไม่มีข้อมูลเฉพาะพื้นที่ (เช่น วันเก็บขยะ จุดรับ) ให้แนะนำติดต่อหน่วยงานนี้ ห้ามเดาข้อมูลเอง")
        return hits, vision, system, msgs, None

    def answer(self, question: str, history=None, image_bytes: bytes | None = None, k: int = 4) -> dict:
        hits, vision, system, msgs, immediate = self._prepare(question, history, image_bytes, k)
        text = immediate if immediate is not None else self._complete(system, msgs, 1200)
        return {"text": text, "sources": hits, "vision": vision}

    def answer_stream(self, question: str, history=None, image_bytes: bytes | None = None, k: int = 4):
        """ตอบทีละส่วน: event แรกคือ {'type': 'meta' หรือ 'full', 'sources', 'vision', ('text' เมื่อ full)}
        ตามด้วย {'type': 'delta', 'text'} จนจบ"""
        hits, vision, system, msgs, immediate = self._prepare(question, history, image_bytes, k)
        if immediate is not None:
            yield {"type": "full", "text": immediate, "sources": hits, "vision": vision}
            return
        yield {"type": "meta", "sources": hits, "vision": vision}
        for piece in self._stream(system, msgs, 1200):
            yield {"type": "delta", "text": piece}

    def _stream(self, system: str, messages: list, max_tokens: int):
        if self.provider == "openai":
            stream = self.client.chat.completions.create(
                model=self.model, max_completion_tokens=max_tokens, stream=True,
                messages=[{"role": "system", "content": system}] + messages)
            for chunk in stream:
                if chunk.choices and chunk.choices[0].delta.content:
                    yield chunk.choices[0].delta.content
        else:
            with self.client.messages.stream(model=self.model, max_tokens=max_tokens,
                                             system=system, messages=messages) as stream:
                yield from stream.text_stream
